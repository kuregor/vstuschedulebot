"""Открытие книги расписания независимо от формата: .xls или .xlsx.

Учебный отдел выкладывает оба: часть факультетов ведёт расписание в Excel
97-2003 (.xls), часть — в современном (.xlsx). Парсер шахматки написан на
интерфейсе листа xlrd (nrows / ncols / cell_value / merged_cells), поэтому
книгу .xlsx мы приводим к тому же интерфейсу, а не переписываем разбор.
"""
from __future__ import annotations

ZIP_MAGIC = b"PK\x03\x04"
OLE_MAGIC = b"\xd0\xcf\x11\xe0"


class SheetLike:
    """Лист в интерфейсе xlrd: значения и объединения, всё уже в памяти.

    Значения читаются сразу целиком: парсер ходит по сетке много раз, а
    поячеечный доступ в openpyxl заметно дороже одного прохода.
    """

    def __init__(self, rows: list[list], merged: list[tuple[int, int, int, int]]):
        self._rows = rows
        self.nrows = len(rows)
        self.ncols = max((len(r) for r in rows), default=0)
        # (r0, r1, c0, c1), правая граница не включается — как в xlrd
        self.merged_cells = merged

    def cell_value(self, r: int, c: int):
        if r < 0 or r >= self.nrows:
            return ""
        row = self._rows[r]
        if c < 0 or c >= len(row):
            return ""
        return row[c]


def _sheet_from_xlsx(path: str) -> SheetLike:
    import openpyxl

    book = openpyxl.load_workbook(path, data_only=True)
    ws = book.worksheets[0]

    rows: list[list] = []
    for row in ws.iter_rows(values_only=True):
        rows.append(["" if v is None else v for v in row])

    merged: list[tuple[int, int, int, int]] = []
    for rng in ws.merged_cells.ranges:
        # openpyxl нумерует с единицы и включает правую границу
        merged.append((rng.min_row - 1, rng.max_row, rng.min_col - 1, rng.max_col))

    book.close()
    return SheetLike(rows, merged)


def _sheet_from_xls(path: str):
    import xlrd

    book = xlrd.open_workbook(path, formatting_info=True)
    return book.sheet_by_index(0)


def detect_format(path: str) -> str:
    """-> 'xls' | 'xlsx'. Смотрим на сигнатуру файла, а не на расширение:
    на сайте попадаются .xls, внутри которых лежит книга нового формата."""
    with open(path, "rb") as fh:
        head = fh.read(8)
    if head.startswith(ZIP_MAGIC):
        return "xlsx"
    if head.startswith(OLE_MAGIC):
        return "xls"
    raise ValueError(
        "Файл не похож на книгу Excel (ожидался .xls или .xlsx). "
        "Скорее всего, по ссылке отдали страницу с ошибкой, а не расписание."
    )


def open_sheet(path: str):
    """Первый лист книги в интерфейсе, который понимает MergedGrid."""
    if detect_format(path) == "xlsx":
        return _sheet_from_xlsx(path)
    return _sheet_from_xls(path)
